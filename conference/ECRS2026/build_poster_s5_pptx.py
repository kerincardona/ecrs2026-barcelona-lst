#!/usr/bin/env python3
"""ECRS 2026 poster for Abstract 1 (Session S5, Remote Sensing in Urban Areas) — v1.

Same official Sciforum/MDPI template, layout and rules as the S1 poster
(build_poster_pptx.py, whose text/table helpers are imported unchanged):
  * every text run >= 24 pt (venue spec), Arial
  * every number traced to a source file (no hard-coded results without provenance)

Abstract : docs/ecrs2026_abstract.md  ("Twelve-Year Landsat-Derived SUHI by LCZ ...")
Paper    : paper/main.tex             (numbers and caveats; paper is newer than the abstract)
Figures  : scripts/figure_poster_s5.py -> fig_s5_maps.png (76 cm wide; (c) = census-section
           exposure map), fig_s5_suhi.png (25 cm, 8 main LCZ classes, middle column)
v1.1 (14-oct-2026): panel (c) SUHI bars -> building-buffer exposure map by census section
  (sits above the "Exposure indicators" column); SUHI by LCZ moved to the middle column
  (8 main classes); SHAP chart dropped, driver result kept as a Conclusions bullet.
  (Census-section aggregation rather than individual footprints: Landsat resolves ~100 m.)
NUMBERS  : fig_s5_provenance.json (written by the figure script from):
  suhi_by_lcz_JJA_daytime.csv          LCZ 8 +1.42 K (316.42 K), LCZ 10 +0.96, LCZ A -8.04, ref D 315.0 K
  hotspot_land_sensitivity.json        P90 317.15 K; >=1 event 64.2 % window / 77.6 % land;
                                       f>0.5 1.07 % / 1.29 %; f_max 0.686; top-decile cut 0.288
  lcz_elevation_confound.json          elevation-matched A - LCZ8 = -8.39 K
  rq3_ml_elevfix_results_land.json     copdem: R2 0.84 +- 0.01, OLS 0.77 +- 0.02, NDVI rank 1 in 5/5,
                                       n = 37,345
  validation_era5.json                 n 325, bias -4.32, r 0.92; JJA n 134, bias -5.57, r 0.26
  validation_ecostress.json            7 scenes, bias +0.37, RMSE 4.42, r 0.92;
                                       per scene bias -6.44..+4.25, r 0.43..0.96
  building_lst_exposure.parquet        106,203 buildings, 13 municipalities, 10,333 distinct 30 m medians,
                                       r(30 m, 100 m) = 0.989
  censustract_lst_exposure_clean.parquet  1,068 sections, median 34.9-47.2 C, IQR 41.1-42.1 C
  paper/main.tex                       378 scenes (153 JJA), 464 km2, ~10:30 overpass, 60 m cloud buffer

Differences vs the submitted abstract (abstract numbers stand; poster adds the paper's caveats):
  * hotspot shares given for both the window (abstract: 64 %, 1.1 %) and land only (78 %, 1.3 %)
  * ECOSTRESS: pooled +0.37 K is cancellation of per-scene biases (-6.4..+4.3 K); spatial
    agreement (per-scene r) is the supported claim, n = 7 gives no absolute accuracy
  * LCZ A -8.04 K includes a topographic component; elevation-matched comparison added
  * driver model (gradient boosting + SHAP) is in the paper, not in the abstract

Run: ./render_poster_s5.sh   (figures first: .venv/bin/python scripts/figure_poster_s5.py)
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Emu, Pt

import build_poster_pptx as base
from build_poster_pptx import (B, I, NEGCTRL, P, TEAL, WINNER, fill, table, textbox, FONT,
                               MIN_PT, TEMPLATE)

HERE = Path(__file__).resolve().parent
OUT = HERE / "ECRS2026_poster_S5_v1.pptx"
FIG = HERE / "fig_s5_maps.png"
FIG_W_CM = 76.0          # MUST equal W in scripts/figure_poster_s5.py
SUHI_FIG = HERE / "fig_s5_suhi.png"
SUHI_W_CM, SUHI_H_CM = 25.0, 9.8   # MUST equal SUW, SUH in scripts/figure_poster_s5.py
PROV = json.load(open(HERE / "fig_s5_provenance.json"))
MINUS = "\u2212"


def half_up(x, nd=1):
    from decimal import Decimal, ROUND_HALF_UP
    return str(Decimal(str(x)).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP))


def k(x, nd=2, sign=True):
    s = half_up(x, nd)
    if sign and not s.startswith("-"):
        s = "+" + s
    return s.replace("-", MINUS)


def _k_unused(x, nd=2, sign=True):
    s = f"{x:+.{nd}f}" if sign else f"{x:.{nd}f}"
    return s.replace("-", MINUS)


def main():
    pv = PROV
    hs, era, eco = pv["hotspot"], pv["era5"], pv["ecostress"]
    dm, bl, ct = pv["driver_model"], pv["buildings"], pv["census"]
    suhi = pv["suhi_K"]

    prs = Presentation(str(TEMPLATE))
    slide = prs.slides[0]
    S = {s.name: s for s in slide.shapes}

    # ================= TITLE BLOCK =============================================
    tb = S["Text Box 40"]
    specs = [
        ("Twelve-Year Landsat-Derived Surface Urban Heat Island Intensity by Local "
         "Climate Zone in Barcelona (2013\u20132024)", 48),
        ("Kerin Cardona*, Jordi Cipriano, Gerard Mor", 32),
        ("International Centre for Numerical Methods in Engineering (CIMNE), "
         "Building Energy and Environment Group, Lleida, Spain", 24),
        ("*Corresponding author: kcardona@cimne.upc.edu", 24),
    ]
    for p, (text, sz) in zip(tb.text_frame.paragraphs, specs):
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].font.size = Pt(sz)
        if sz <= 24:
            p.runs[0].font.name = FONT
        p.space_after = Pt(6)
        base._used_sizes.append(("title", sz))
    tb.top = Emu(3573659)

    # ================= INTRODUCTION & AIM =====================================
    fill(S["TextBox 20"], [
        P("Outdoor thermal exposure depends strongly on urban surface temperature, yet "
          "long-term, high-resolution land surface temperature (LST) climatologies remain "
          "scarce for Mediterranean cities:", size=30),
        P(B("Weather stations: "), "point air temperature only", bullet=True, size=30),
        P(B("ERA5-Land: "), "skin temperature at ~9 km, no intra-urban structure",
          bullet=True, size=30),
        P(B("Single satellite scenes: "), "snapshots that depend on the weather of the day",
          bullet=True, size=30),
        P(B("Aim: "), "a 12-year, 30 m LST climatology for Barcelona, read through Local Climate "
          "Zones (LCZ) and turned into hotspot, building and census-section exposure indicators "
          "for adaptation planning.", size=30, space_after=0),
    ], space_after=12)

    # ================= METHOD ==================================================
    fill(S["TextBox 24"], [
        P("Barcelona metropolitan core: 464 km\u00b2 window (~21 \u00d7 22 km), "
          "13 municipalities.", size=30),
        P(B("378 "), "cloud-filtered Landsat 8/9 Collection 2 Level-2 scenes, 2013\u20132024 "
          "(153 in summer), ~10:30 local; 30 m grid, ~100 m native thermal", bullet=True, size=30),
        P(B("Climatology: "), "seasonal median, MAD and clear-sky count per pixel",
          bullet=True, size=30),
        P(B("SUHI by LCZ "), "(Bechtel et al. 2019): class median minus LCZ D "
          f"(low plants, {pv['lczD_ref_K']:.1f} K)", bullet=True, size=30),
        P(B("Hotspot frequency: "), "share of summer scenes above the 90th percentile of "
          f"the JJA field ({half_up(pv['p90_threshold_K'])} K)", bullet=True, size=30),
        P(B("Drivers: "), f"gradient boosting + SHAP, 13 covariates, {dm['n']:,} land cells, "
          "2 km spatial-block cross-validation", bullet=True, size=30),
        P(B("Cross-checks: "), f"ERA5-Land ({era['n']} pairs); ECOSTRESS 70 m "
          f"({eco['n_scenes']} scenes, |\u0394t| \u2264 2 h)", bullet=True, size=30,
          space_after=0),
    ], space_after=10)

    # ================= RESULTS & DISCUSSION ===================================
    res = S["TextBox 31"]
    L, T, W = res.left, res.top, res.width
    fill(res, [P(B("Summer (JJA) daytime, 12-year climatology. "),
                 f"{bl['n']:,} buildings and {ct['n']:,} census sections linked to it.",
                 size=26)])

    y = T + Cm(1.7)
    fig_w = Cm(FIG_W_CM)
    fx = L + (W - fig_w) / 2
    pic = slide.shapes.add_picture(str(FIG), Emu(int(fx)), Emu(int(y)), width=Emu(int(fig_w)))
    cap_top = y + pic.height + Cm(0.2)
    textbox(slide, fx, cap_top, fig_w, Cm(2.3), [
        P(B("(a) "), "The Llobregat industrial corridor and dense cores exceed 45 \u00b0C; "
          "Collserola forest is the main cool feature. ",
          B("(b) "), f"{hs['window_ge1_pct']:.0f} % of the window ({hs['land_ge1_pct']:.0f} % "
          f"of land) is a hotspot at least once, {hs['window_gt05_pct']:.1f} % "
          f"({hs['land_gt05_pct']:.1f} %) in over half of summer scenes. ",
          B("(c) "), f"Median LST in 30 m building buffers, {ct['n']:,} census sections of "
          "Barcelona city; port and logistics areas are the hottest.", size=24, space_after=0)])

    # --- three columns ------------------------------------------------------
    cy = cap_top + Cm(2.6)
    gap = Cm(1.0)
    cw = (W - 2 * gap) / 3
    cw_cm = cw / 360000
    c1, c2, c3 = L, L + cw + gap, L + 2 * (cw + gap)

    textbox(slide, c1, cy, cw, Cm(1.3), [
        P(B("Cross-checks of the 30 m field", TEAL), size=26, space_after=0)])
    sb, sr = eco["scene_bias_range"], eco["scene_r_range"]
    table(slide, c1, cy + Cm(1.5), [cw_cm * 0.42, cw_cm * 0.15, cw_cm * 0.25, cw_cm * 0.18], 1.15,
          [["Reference", "n", "Bias (K)", "r"],
           ["ERA5-Land, all", str(era["n"]), k(era["bias"]), f"{era['r']:.2f}"],
           ["ERA5-Land, JJA", str(era["by_season"]["JJA"]["n"]),
            k(era["by_season"]["JJA"]["bias"]), f"{era['by_season']['JJA']['r']:.2f}"],
           ["ECOSTRESS, pooled", str(eco["n_scenes"]), k(eco["bias"]), f"{eco['r']:.2f}"],
           ["ECOSTRESS, per scene", str(eco["n_scenes"]),
            f"{k(sb[0], 1)}\u2026{k(sb[1], 1)}", f"{sr[0]:.2f}\u2013{sr[1]:.2f}"]],
          row_colors={2: NEGCTRL})
    textbox(slide, c1, cy + Cm(1.5 + 5 * 1.15 + 0.3), cw, Cm(3.5), [
        P(I("Bias = reference \u2212 Landsat; n = matched Landsat scenes. ERA5 follows the "
            "seasons but not the city. ECOSTRESS places warm and cool areas in the same "
            "locations; n = 7 cannot establish absolute accuracy."), size=24, space_after=0)])

    textbox(slide, c2, cy, cw, Cm(1.3), [
        P(B("SUHI by Local Climate Zone (JJA)", TEAL), size=26, space_after=0)])
    slide.shapes.add_picture(str(SUHI_FIG), Emu(int(c2)), Emu(int(cy + Cm(1.3))),
                             width=Emu(int(Cm(SUHI_W_CM))))
    em = pv["elev_matched_A_minus"]["LCZ8"]
    textbox(slide, c3, cy, cw, Cm(11), [
        P(B("Exposure indicators", TEAL), size=26, space_after=4),
        P(f"{bl['n']:,} buildings, {bl['municipalities']} municipalities: LST in 30 m and "
          f"100 m buffers (r = {bl['r_30_vs_100']:.3f})", bullet=True, size=24, space_after=6),
        P(f"Only {bl['distinct_30m_medians']:,} distinct values: Landsat resolves ~100 m "
          "context, not single buildings", bullet=True, size=24, space_after=6),
        P(f"{ct['n']:,} census sections: median {ct['median_C_min_max'][0]:.1f}\u2013"
          f"{ct['median_C_min_max'][1]:.1f} \u00b0C, but IQR only {ct['median_C_iqr'][0]:.1f}"
          f"\u2060\u2013\u2060{ct['median_C_iqr'][1]:.1f} \u00b0C; extremes set by land use (port vs "
          "seafront, parks)", bullet=True, size=24, space_after=0),
    ])
    print(f"[layout] fig {y/360000:.1f}-{(y+pic.height)/360000:.1f} cm, columns top "
          f"{cy/360000:.1f} cm, SUHI column bottom {(cy + Cm(1.3 + SUHI_H_CM))/360000:.1f} cm "
          f"(results group ends 99.2 cm, Conclusions header 101.1 cm)")

    # ================= CONCLUSIONS ============================================
    fill(S["TextBox 42"], [
        P(f"A tree-canopy cool island ({k(suhi['A'], 1)} K; still {abs(em):.1f} K below LCZ 8 "
          f"at equal elevation) dominates over a weak built heat island "
          f"(\u2264 {k(suhi['8'], 1)} K).", bullet=True, size=24, space_after=4),
        P(f"Vegetation (NDVI) is the leading driver (held-out R\u00b2 = {dm['r2_mean_sd'][0]:.2f}; "
          f"ranked first in {dm['ndvi_rank_per_seed'].count(1)}/5 splits): tree cover is the "
          "main modifiable cooling lever.",
          bullet=True, size=24, space_after=4),
        P("Reproducible 30 m baseline from open data for adaptation planning; ERA5-Land "
          "cannot replace it.", bullet=True, size=24, space_after=0),
    ])

    # ================= FUTURE WORK / REFERENCES ===============================
    fill(S["TextBox 43"], [
        P(B("Future work: "), "combine exposure with socio-economic vulnerability; "
          "year-split validation; night-time LST.", size=24, space_after=4),
        P(B("Data: "), "Landsat 8/9 C2 L2, ECOSTRESS L2G v002, ERA5-Land, INSPIRE cadastre, "
          "Copernicus DEM / HRL / Urban Atlas; LCZ map: Demuzere et al. (2019).",
          size=24, space_after=4),
        P(B("Session S5 \u2014 Remote Sensing in Urban Areas"), size=24, space_after=0),
    ])

    prs.save(str(OUT))
    sizes = base._used_sizes
    bad = [(n, s) for n, s in sizes if s < MIN_PT]
    print(f"[fonts] runs checked: {len(sizes)} | min = {min(s for _, s in sizes)} pt"
          f" | below {MIN_PT} pt: {bad if bad else 'none'}")
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
