#!/usr/bin/env python3
"""ECRS 2026 poster (Abstract 2, Session S1) on the OFFICIAL Sciforum/MDPI template.

v2 (12-oct-2026) — second review (v1 of this script is in git: commit 4f9e0f6):
  * every text run >= 24 pt (venue spec: "Minimum font size: >=24")
  * results block re-laid out to use its full area (figure ~52 cm wide)
  * content corrected against the source JSONs (see NUMBERS below)
  * acknowledgment removed (no funding info provided; do not invent)

Template : templates/2026-Poster-Template.pptx (official Sciforum template, not redistributed)
Output   : ECRS2026_poster_S1_official_template.pptx (+ .pdf via LibreOffice)

NUMBERS — source of every figure on the poster
  data/interim/paired_s3_l8/baselines/baseline_results.json  (table 1)
  data/interim/paired_s3_l8/baseline3_vs_ecostress.json      (table 2)
  data/interim/paired_s3_l8/phase_a2_xgboost_anomaly_results.json
      XGB raw S3 rho 0.6049 (ties B2 0.6004); XGB anomaly rho 0.7013,
      RMSE 4.2463, hotspot F1 0.3205, IoU 0.2815
  data/interim/paired_s3_l8/verify_residual_ceiling_results.json
      fraction of residual variance explainable by static features 0.0102
  docs/viac_downscaling_gonogo.md
      B3 better on 38/61 dates, one-sided Wilcoxon p=0.0033 (naive LCZ/NDVI; superseded
      by the georef rerun below:
      wilcoxon_baseline_vs_a2_results.json NOT present locally)
  fig_poster_example_provenance.json  (per-scene rho 0.20 / 0.92; 3 K vs 21 K)

Corrections vs poster_v1.tex (not supported by the JSONs above):
  "raw S3 underperforms B2"  -> ties B2 (0.605 vs 0.600)
  "closes 78 % of the gap"   -> arithmetic gives 69.5 %; claim dropped
  "still loses to B3"        -> only on rho/IoU; lower RMSE and higher F1 now stated
  "< 1 % residual variance"  -> 1.02 %, now "~1 %"

v4 (13-oct-2026): ALL Sentinel-3 numbers replaced by the DAY-PASS rerun (the original
  download used the ~21 UTC night pass; see docs/viac_downscaling_gonogo.md "REDESCARGA CON
  PASADA DIURNA"). Sources: data/interim/paired_s3_l8/baselines_daypass/
  {baseline_results, baseline3_vs_ecostress, phase_a_xgboost_results,
  phase_a2_xgboost_anomaly_results, wilcoxon_baseline_vs_a2_results,
  verify_residual_ceiling_results}.json; figure fig_poster_example_4p_provenance.json
  (S3 rho 0.90, B2 0.83, B3 0.97; S3 range 10.9 K vs Landsat 21.0 K). 189 pairs; split 80/34/75.
  Abstract correction note added in the footer box. The v1-v3 provenance below is superseded.
v3 (13-oct-2026): 4-panel figure (S3 / B2 / B3 / Landsat), ECOSTRESS moved below.
  B2 row now uses the GEOREFERENCED LCZ (scripts/qa_b2_lcz_registration.py and evaluate_baselines.py ->
  data/interim/paired_s3_l8/baselines/qa_b2_lcz_registration.json): the published
  evaluate_baselines.load_lcz() resized the EPSG:4326 LCZ onto the EPSG:32631 grid with
  scipy zoom (no reprojection); class agreement 0.62; rho 0.600 -> 0.651 (Wilcoxon
  p=1e-11); IoU/F1/RMSE essentially unchanged.
  Static-covariate bullet RESTORED after rerun with georeferenced LCZ+NDVI
  (verify_residual_ceiling.py, now georef by default -> verify_residual_ceiling_results.json):
  [A] within-pixel share 1 - 0.0102 (independent of features by construction);
  [D] M1 static residual model test R2 0.0033, rho 0.7436 -> 0.7459 (naive run: 0.0032,
  0.7457). Registration does not change the conclusion.
  Learned-model bullets now from the georeferenced runs. Since 13-oct the canonical
  scripts load LCZ/NDVI via scripts/grid_align.py, so the canonical JSONs ARE the georef
  results; the old scipy-zoom outputs are kept as *_naive_zoom.json:
    phase A raw-S3 XGB   rho 0.6049 -> 0.6341 (still < B2 georef 0.6514)
    phase A2 anomaly XGB rho 0.7013 -> 0.7027, RMSE 4.2463 -> 4.2411, F1 0.3205 -> 0.3260
    Wilcoxon B3 vs A2    wins 38/61 -> 39/61, one-sided p 0.0033 -> 0.0041

"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Emu, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "templates" / "2026-Poster-Template.pptx"  # official template, see README
OUT = HERE / "ECRS2026_poster_S1_official_template.pptx"
FIG = HERE / "fig_poster_example_4p.png"  # scripts/figure_poster_example.py --with-b2 --print-width-cm FIG_W_CM
FIG_W_CM = 70.0   # MUST equal --print-width-cm in make_fig_4p.sh (fonts are sized for it)

FONT = "Arial"
MIN_PT = 24

TEAL = RGBColor(0x12, 0x5F, 0x55)       # template header colour
WINNER = RGBColor(0x1E, 0x84, 0x49)
NEGCTRL = RGBColor(0xB0, 0x3A, 0x2E)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
B3_FILL = RGBColor(0xE4, 0xF3, 0xEC)
B1_FILL = RGBColor(0xF8, 0xE5, 0xE3)
ZEBRA = RGBColor(0xF2, 0xF2, 0xF2)

_used_sizes = []


# --------------------------------------------------------------------------
# text helpers
# --------------------------------------------------------------------------
def _bullet(para, on, indent_cm=0.75):
    pPr = para._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if on:
        pPr.set("marL", str(int(Cm(indent_cm))))
        pPr.set("indent", str(-int(Cm(indent_cm))))
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "\u2022")
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))


def fill(shape, paras, size=24, align=PP_ALIGN.LEFT, space_after=8):
    """paras: list of dicts {runs:[(text, bold, italic, color)], bullet:bool, size:int}"""
    tf = shape.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, spec in enumerate(paras):
        para = first if i == 0 else tf.add_paragraph()
        para.alignment = spec.get("align", align)
        para.space_after = Pt(spec.get("space_after", space_after))
        _bullet(para, spec.get("bullet", False))
        sz = spec.get("size", size)
        for text, bold, italic, color in spec["runs"]:
            run = para.add_run()
            run.text = text
            run.font.size = Pt(sz)
            run.font.bold = bold
            run.font.italic = italic
            run.font.name = FONT
            run.font.color.rgb = color or BLACK
            _used_sizes.append((shape.name, sz))


def P(*runs, bullet=False, size=24, **kw):
    """Paragraph spec. runs: str (plain) or tuple (text, bold, italic, color)."""
    norm = []
    for r in runs:
        if isinstance(r, str):
            norm.append((r, False, False, None))
        else:
            r = tuple(r) + (False, False, None)[len(r) - 1:]
            norm.append(r[:4])
    return dict(runs=norm, bullet=bullet, size=size, **kw)


def B(text, color=None):
    return (text, True, False, color)


def I(text):
    return (text, False, True, None)


def textbox(slide, left, top, width, height, paras, **kw):
    tb = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    tb.text_frame.margin_left = tb.text_frame.margin_right = Cm(0.15)
    fill(tb, paras, **kw)
    return tb


def table(slide, left, top, col_widths_cm, row_h_cm, rows, size=24,
          header_fill=TEAL, row_fills=None, row_colors=None, bold_rows=()):
    n_rows, n_cols = len(rows), len(rows[0])
    width = sum(Cm(w) for w in col_widths_cm)
    gs = slide.shapes.add_table(n_rows, n_cols, Emu(int(left)), Emu(int(top)),
                                Emu(int(width)), Emu(int(Cm(row_h_cm) * n_rows)))
    t = gs.table
    tblPr = gs._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    for j, w in enumerate(col_widths_cm):
        t.columns[j].width = Cm(w)
    for i in range(n_rows):
        t.rows[i].height = Cm(row_h_cm)
        for j in range(n_cols):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Cm(0.25)
            c.margin_top = c.margin_bottom = Cm(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = rows[i][j]
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            run = para.runs[0]
            run.font.size = Pt(size)
            run.font.name = FONT
            _used_sizes.append(("table", size))
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = header_fill
                run.font.bold = True
                run.font.color.rgb = WHITE
            else:
                fillc = (row_fills or {}).get(i, WHITE if i % 2 else ZEBRA)
                c.fill.fore_color.rgb = fillc
                run.font.color.rgb = (row_colors or {}).get(i, BLACK)
                run.font.bold = i in bold_rows
    return gs


# --------------------------------------------------------------------------
def main():
    prs = Presentation(str(TEMPLATE))
    slide = prs.slides[0]
    S = {s.name: s for s in slide.shapes}

    # ================= TITLE BLOCK (fixed gap above divider "Line 21") =========
    tb = S["Text Box 40"]
    ps = tb.text_frame.paragraphs
    specs = [
        ("Towards a Spatio-Temporally Complete Urban Land Surface Temperature Record: "
         "A Multi-Sensor Benchmark and Rank-Based Baseline Evaluation for Thermal "
         "Downscaling over Barcelona", 48),
        ("Kerin Cardona*, Jordi Cipriano, Gerard Mor", 32),
        ("International Centre for Numerical Methods in Engineering (CIMNE), "
         "Building Energy and Environment Group, Lleida, Spain", 24),
        ("*Corresponding author: kcardona@cimne.upc.edu", 24),
    ]
    for p, (text, sz) in zip(ps, specs):
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].font.size = Pt(sz)
        if sz <= 24:   # affiliation/contact: template font is not installed and its
            p.runs[0].font.name = FONT   # fallback renders "kcardona@ cimne" with a gap
        p.space_after = Pt(6)
        _used_sizes.append(("title", sz))
    tb.top = Emu(3573659)

    # ================= INTRODUCTION & AIM =====================================
    fill(S["TextBox 20"], [
        P("Urban heat-exposure mapping needs high-resolution land surface temperature "
          "(LST), yet no single sensor is complete in both space and time:", size=28),
        P(B("Landsat 8/9: "), "30 m grid (100 m native thermal), 16-day revisit, "
          "cloud-limited", bullet=True, size=28),
        P(B("Sentinel-3 SLSTR: "), "daily, but 1 km", bullet=True, size=28),
        P(B("ECOSTRESS: "), "70 m, irregular revisit", bullet=True, size=28),
        P("Standardised, exposure-relevant benchmarks for urban thermal downscaling "
          "are lacking.", bullet=True, size=28),
        P(B("Aim: "), "build a multi-sensor benchmark and set a verified baseline floor "
          "that any learned downscaling model must beat.", size=28, space_after=0),
    ], space_after=12)

    # ================= METHOD ==================================================
    fill(S["TextBox 24"], [
        P("Benchmark over the Barcelona metropolitan core (464 km\u00b2, "
          "~21 \u00d7 22 km):", size=28),
        P(B("189 "), "co-registered pairs, 2016\u20132024: Sentinel-3 SLSTR S8 brightness "
          "temperature (1 km, day pass \u224810 UTC) / Landsat LST (30 m)", bullet=True, size=28),
        P("Independent ECOSTRESS (70 m) reference, never used for fitting",
          bullet=True, size=28),
        P(B("Temporal split, no leakage: "), "train 2016\u20132021 (80) \u00b7 validation "
          "2022 (34) \u00b7 held-out test 2023\u20132024 (75)", bullet=True, size=28),
        P(B("Rank-based protocol: "), "Spearman \u03c1, hottest-10 % IoU, hotspot F1, "
          "per-LCZ consistency (exposure depends on spatial ranking, not magnitude)",
          bullet=True, size=28),
        P(B("Baselines: "), "B1 bilinear resample of S3 (no downscaling) \u00b7 "
          "B2 mean LST per Local Climate Zone \u00b7 B3 = Landsat train climatology + "
          "daily S3 anomaly", bullet=True, size=28, space_after=0),
    ], space_after=12)

    # ================= RESULTS & DISCUSSION ===================================
    res = S["TextBox 31"]           # keep template position (built-in header margin)
    L, T, W = res.left, res.top, res.width
    fill(res, [P(B("Held-out test, 75 pairs (2023\u20132024), mean \u00b1 s.d. over dates. "),
                 "B1 = S3 without downscaling (reference).", size=26)])

    y = T + Cm(1.7)
    rows1 = [
        ["Baseline", "Spearman \u03c1 \u2191", "IoU top-10 % \u2191", "Hotspot F1 \u2191",
         "RMSE (K) \u2193", "Bias (K)"],
        ["B1  S3 bilinear resample (no downscaling)", "0.73 \u00b1 0.14", "0.26 \u00b1 0.13",
         "0.00 \u00b1 0.00", "7.10 \u00b1 3.73", "\u22126.42"],
        ["B2  Mean LST per LCZ", "0.62 \u00b1 0.20", "0.04 \u00b1 0.02", "0.09 \u00b1 0.10",
         "10.22 \u00b1 4.77", "\u22120.20"],
        ["B3  Landsat climatology + S3 anomaly", "0.79 \u00b1 0.21", "0.37 \u00b1 0.14",
         "0.33 \u00b1 0.17", "3.98 \u00b1 1.83", "+1.36"],
    ]
    col_w = [26.0, 10.0, 10.5, 10.0, 10.5, 8.7]       # = 75.7 cm < 77.3 cm box
    table(slide, L + Cm(0.2), y, col_w, 1.35, rows1,
          row_fills={1: B1_FILL, 3: B3_FILL},
          row_colors={1: NEGCTRL, 3: WINNER}, bold_rows=(3,))
    y += Cm(1.35 * 4 + 0.25)
    textbox(slide, L, y, W, Cm(1.2), [
        P(I("Chance level for top-10 % IoU \u2248 0.05. B2 does not locate hotspots; B1 ranks "
            "well but its top-of-atmosphere values miss every absolute hotspot (F1 = 0)."),
          size=24, space_after=0)])
    y += Cm(1.9)

    # --- full-width 4-panel worked example -----------------------------------
    fig_w = Cm(FIG_W_CM)
    fx = L + (W - fig_w) / 2
    pic = slide.shapes.add_picture(str(FIG), Emu(int(fx)), Emu(int(y)), width=Emu(int(fig_w)))
    cap_top = y + pic.height + Cm(0.2)
    textbox(slide, fx, cap_top, fig_w, Cm(2.4), [
        P(B("Held-out date 2024-08-03, per-scene \u03c1. "),
          "(a) At 1 km S3 already ranks the scene but spans 11 K vs 21 K, with no 30 m detail. "
          "(b) B2 has the LCZ pattern but no daily signal (all-season mean, own scale). "
          "(c) B3 recovers the intra-urban detail. (d) Reference.", size=24, space_after=0)])

    # --- three columns below the figure -------------------------------------
    cy = cap_top + Cm(2.9)
    gap = Cm(1.0)
    cw = (W - 2 * gap) / 3
    c1, c2, c3 = L, L + cw + gap, L + 2 * (cw + gap)

    textbox(slide, c1, cy, cw, Cm(2.4), [
        P(B("Cross-check: ECOSTRESS (70 m)", TEAL), size=26, space_after=2),
        P("3 test dates with usable swath", size=24, space_after=0)])
    cw_cm = cw / 360000
    table(slide, c1, cy + Cm(2.6), [cw_cm * 0.46, cw_cm * 0.27, cw_cm * 0.27], 1.2,
          [["Comparison", "\u03c1", "RMSE (K)"],
           ["B3 vs ECOSTRESS", "0.79 \u00b1 0.11", "4.52 \u00b1 1.23"],
           ["ECOSTRESS vs Landsat", "0.91 \u00b1 0.03", "\u2014"]])
    textbox(slide, c1, cy + Cm(2.6 + 3.6 + 0.3), cw, Cm(2.2), [
        P(I("n = 3: a consistency check, too small to rank methods."), size=24, space_after=0)])

    textbox(slide, c2, cy, cw, Cm(11), [
        P(B("Why not a learned model (yet)?", TEAL), size=26, space_after=4),
        P("Gradient boosting on raw S3: \u03c1 = 0.78, above the LCZ mean (B2, 0.62), "
          "close to B3 (0.79).",
          bullet=True, size=24, space_after=6),
        P("With S3 anomaly + climatology: \u03c1 = 0.81 vs 0.79, not significant (B3 wins "
          "41/75 dates, p = 0.71), with lower RMSE (3.01 vs 3.98 K) and higher F1 (0.43 vs 0.33).",
          bullet=True, size=24, space_after=6),
        P("Static covariates (NDVI, elevation, building density, LCZ) as a residual model "
          "on B3: test R\u00b2 \u2248 0, \u03c1 0.79 \u2192 0.80. \u2248 97 % of the "
          "residual variance is within-pixel, i.e. temporal.",
          bullet=True, size=24, space_after=0),
    ])

    textbox(slide, c3, cy, cw, Cm(11), [
        P(B("B3 today: a climatology-anchored gap-filler, not a downscaler", TEAL),
          size=26, space_after=4),
        P(B("OK  ", WINNER), "seasonal ranking of exposure between districts (\u03c1 = 0.79)",
          size=24, space_after=3),
        P(B("OK  ", WINNER), "filling cloudy or non-overpass days", size=24, space_after=3),
        P(B("NO  ", NEGCTRL), "day-specific hotspot alerts (F1 = 0.33)", size=24, space_after=3),
        P(B("NO  ", NEGCTRL), "absolute daily temperature (RMSE 4.0 K)", size=24, space_after=0),
    ])
    print(f"[layout] fig top {y/360000:.1f} cm, fig bottom {(y+pic.height)/360000:.1f} cm, "
          f"columns top {cy/360000:.1f} cm (conclusions header at {36489997/360000:.1f} cm)")

    # ================= CONCLUSIONS ============================================
    fill(S["TextBox 42"], [
        P("A high-resolution climatological prior is indispensable: learned models "
          "must beat it to add value.", bullet=True, size=24, space_after=4),
        P("The B3 residual is temporal: progress needs temporal information, not more "
          "static covariates.", bullet=True, size=24, space_after=4),
        P("Benchmark, protocol and baselines: a reproducible foundation for "
          "spatio-temporal deep learning.", bullet=True, size=24, space_after=0),
    ])

    # ================= FUTURE WORK / REFERENCES ===============================
    fill(S["TextBox 43"], [
        P(B("Future work: "), "evaluate spatio-temporal deep-learning models on this "
          "benchmark, against B3.", size=24, space_after=4),
        P(B("Data: "), "Sentinel-3 SLSTR, Landsat 8/9 Collection 2, ECOSTRESS L2G LSTE; "
          "LCZ map: WUDAPT (Demuzere et al.).", size=24, space_after=4),
        P(I("Correction to the abstract: the Sentinel-3 overpass selection was fixed after "
            "submission (night \u2192 day pass); all results shown use the day pass."),
          size=24, space_after=4),
        P(B("Session S1 \u2014 AI in Remote Sensing"), size=24, space_after=0),
    ])

    prs.save(str(OUT))
    bad = [(n, s) for n, s in _used_sizes if s < MIN_PT]
    print(f"[fonts] runs checked: {len(_used_sizes)} | min = {min(s for _, s in _used_sizes)} pt"
          f" | below {MIN_PT} pt: {bad if bad else 'none'}")
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
